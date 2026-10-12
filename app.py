import os
import re
import io
import hmac
import sqlite3
import zipfile
from datetime import datetime
from functools import wraps
from pathlib import Path

from flask import (
    Flask, render_template, request, redirect, url_for, flash,
    session, send_file, abort
)
from werkzeug.utils import secure_filename
from werkzeug.security import check_password_hash, generate_password_hash
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from PIL import Image, UnidentifiedImageError

APP_DIR = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get("DATA_DIR", str(APP_DIR / "data")))
UPLOAD_DIR = DATA_DIR / "uploads"
DB_PATH = Path(os.environ.get("DB_PATH", str(DATA_DIR / "students.sqlite3")))
DATA_DIR.mkdir(parents=True, exist_ok=True)
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "CHANGE-THIS-TO-A-LONG-RANDOM-SECRET")
app.config["MAX_CONTENT_LENGTH"] = 8 * 1024 * 1024  # har bir so'rov uchun 8 MB
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = os.environ.get("COOKIE_SECURE", "0") == "1"

ALLOWED_EXTENSIONS = {"jpg", "jpeg", "png", "webp"}
MAX_IMAGE_BYTES = 7 * 1024 * 1024
MAX_IMAGE_SIDE = 5000

# Guruhlarni Render/GitHub sozlamalarida GROUPS=101,102,103 ko'rinishida o'zgartiring.
GROUPS = [
    g.strip() for g in os.environ.get(
        "GROUPS", "101,102,103,104,105,106,107,108,109,110,111,112,113,114,115,116,117,118,201,202,203,204,205,206,207,208,209,210,211,212,213"
    ).split(",") if g.strip()
]

def db():
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    return con

def init_db():
    with db() as con:
        con.execute("""
            CREATE TABLE IF NOT EXISTS students (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                surname TEXT NOT NULL,
                first_name TEXT NOT NULL,
                patronymic TEXT NOT NULL,
                phone TEXT NOT NULL,
                group_name TEXT NOT NULL,
                image_filename TEXT NOT NULL,
                created_at TEXT NOT NULL
            )
        """)

init_db()

def admin_password_ok(password):
    configured_hash = os.environ.get("ADMIN_PASSWORD_HASH", "").strip()
    configured_password = os.environ.get("ADMIN_PASSWORD", "").strip()
    if configured_hash:
        try:
            return check_password_hash(configured_hash, password)
        except ValueError:
            return False
    # Oddiy env parol taqqoslanadi; kuchli parol ishlating.
    return bool(configured_password) and hmac.compare_digest(password, configured_password)

def admin_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get("admin_logged_in"):
            return redirect(url_for("admin_login", next=request.path))
        return fn(*args, **kwargs)
    return wrapper

def clean_text(value, max_len=100):
    value = re.sub(r"\s+", " ", (value or "").strip())
    if not value or len(value) > max_len:
        return None
    if any(ord(ch) < 32 for ch in value):
        return None
    return value

def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS

def safe_image_save(file_storage, group_name, surname, first_name):
    if not file_storage or not file_storage.filename:
        raise ValueError("Rasm tanlanmadi.")
    original = secure_filename(file_storage.filename)
    if not original or not allowed_file(original):
        raise ValueError("Rasm JPG, JPEG, PNG yoki WEBP formatida bo‘lishi kerak.")

    raw = file_storage.read(MAX_IMAGE_BYTES + 1)
    if not raw:
        raise ValueError("Tanlangan rasm bo‘sh.")
    if len(raw) > MAX_IMAGE_BYTES:
        raise ValueError("Rasm hajmi 7 MB dan oshmasligi kerak.")

    try:
        image = Image.open(io.BytesIO(raw))
        image.verify()
        image = Image.open(io.BytesIO(raw))
        if image.format not in {"JPEG", "PNG", "WEBP"}:
            raise ValueError("Rasm formati qo‘llab-quvvatlanmaydi.")
        if image.width > MAX_IMAGE_SIDE or image.height > MAX_IMAGE_SIDE:
            raise ValueError("Rasm o‘lchami juda katta (maksimum 5000×5000).")
        image = image.convert("RGB")
        out = io.BytesIO()
        image.save(out, format="JPEG", quality=92, optimize=True)
        image_bytes = out.getvalue()
    except (UnidentifiedImageError, OSError):
        raise ValueError("Yuklangan fayl haqiqiy rasm emas.")

    group_dir = UPLOAD_DIR / secure_filename(group_name)
    group_dir.mkdir(parents=True, exist_ok=True)

    # Fayl nomini bir-biriga o'xshash talabalar uchun ham takrorlamaslik.
    stamp = datetime.utcnow().strftime("%Y%m%d%H%M%S%f")
    stem = secure_filename(f"{surname}_{first_name}")[:70] or "talaba"
    filename = f"{stem}_{stamp}.jpg"
    destination = group_dir / filename
    destination.write_bytes(image_bytes)
    return f"{secure_filename(group_name)}/{filename}"

@app.after_request
def security_headers(response):
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    response.headers["Permissions-Policy"] = "camera=(), microphone=(), geolocation=()"
    return response

@app.route("/", methods=["GET", "POST"])
def index():
    if request.method == "POST":
        surname = clean_text(request.form.get("surname"), 80)
        first_name = clean_text(request.form.get("first_name"), 80)
        patronymic = clean_text(request.form.get("patronymic"), 80)
        phone = clean_text(request.form.get("phone"), 30)
        group_name = clean_text(request.form.get("group_name"), 30)

        if not all([surname, first_name, patronymic, phone, group_name]):
            flash("Barcha maydonlarni to‘g‘ri to‘ldiring.", "error")
            return render_template("index.html", groups=GROUPS), 400
        if group_name not in GROUPS:
            flash("Guruh ro‘yxatdan tanlanmagan yoki noto‘g‘ri.", "error")
            return render_template("index.html", groups=GROUPS), 400
        if not re.fullmatch(r"[+()\-\d\s]{7,30}", phone):
            flash("Telefon raqamini tekshiring. Masalan: +998 90 123 45 67", "error")
            return render_template("index.html", groups=GROUPS), 400

        try:
            relative_image = safe_image_save(
                request.files.get("photo"), group_name, surname, first_name
            )
        except ValueError as exc:
            flash(str(exc), "error")
            return render_template("index.html", groups=GROUPS), 400

        try:
            with db() as con:
                con.execute("""
                    INSERT INTO students
                    (surname, first_name, patronymic, phone, group_name, image_filename, created_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                """, (surname, first_name, patronymic, phone, group_name,
                      relative_image, datetime.utcnow().isoformat(timespec="seconds")))
        except Exception:
            # Bazaga yozilmasa, yetim rasmni tozalashga uriniladi.
            try:
                (UPLOAD_DIR / relative_image).unlink(missing_ok=True)
            except OSError:
                pass
            flash("Ma’lumotni saqlashda xatolik yuz berdi. Qayta urinib ko‘ring.", "error")
            return render_template("index.html", groups=GROUPS), 500

        return render_template("index.html", groups=GROUPS, submitted=True)

    return render_template("index.html", groups=GROUPS)

@app.route("/admin", methods=["GET"])
@admin_required
def admin():
    search = (request.args.get("q") or "").strip()[:100]
    group_filter = (request.args.get("group") or "").strip()
    query = "SELECT * FROM students WHERE 1=1"
    params = []
    if search:
        like = f"%{search}%"
        query += " AND (surname LIKE ? OR first_name LIKE ? OR patronymic LIKE ? OR phone LIKE ?)"
        params.extend([like, like, like, like])
    if group_filter in GROUPS:
        query += " AND group_name = ?"
        params.append(group_filter)
    query += " ORDER BY group_name, surname COLLATE NOCASE, first_name COLLATE NOCASE"
    with db() as con:
        students = con.execute(query, params).fetchall()
        total = con.execute("SELECT COUNT(*) FROM students").fetchone()[0]
        group_counts = con.execute(
            "SELECT group_name, COUNT(*) AS n FROM students GROUP BY group_name ORDER BY group_name"
        ).fetchall()
    return render_template("admin.html", students=students, groups=GROUPS,
                           total=total, group_counts=group_counts,
                           search=search, group_filter=group_filter)

@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    if session.get("admin_logged_in"):
        return redirect(url_for("admin"))
    if request.method == "POST":
        password = request.form.get("password", "")
        if admin_password_ok(password):
            session.clear()
            session["admin_logged_in"] = True
            session.permanent = True
            return redirect(url_for("admin"))
        flash("Parol noto‘g‘ri.", "error")
    if not os.environ.get("ADMIN_PASSWORD") and not os.environ.get("ADMIN_PASSWORD_HASH"):
        setup_warning = True
    else:
        setup_warning = False
    return render_template("login.html", setup_warning=setup_warning)

@app.route("/admin/logout", methods=["POST"])
@admin_required
def admin_logout():
    session.clear()
    flash("Tizimdan chiqdingiz.", "success")
    return redirect(url_for("admin_login"))

def make_excel(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Talabalar"
    headers = ["№", "Familiya", "Ism", "Sharif", "Telefon", "Guruh", "Rasm fayli", "Kiritilgan vaqt (UTC)"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for number, row in enumerate(rows, start=1):
        ws.append([number, row["surname"], row["first_name"], row["patronymic"],
                   row["phone"], row["group_name"], row["image_filename"], row["created_at"]])
    widths = [7, 22, 22, 24, 22, 14, 45, 25]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 24
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

@app.route("/admin/export/excel")
@admin_required
def export_excel():
    with db() as con:
        rows = con.execute(
            "SELECT * FROM students ORDER BY group_name, surname COLLATE NOCASE, first_name COLLATE NOCASE"
        ).fetchall()
    output = make_excel(rows)
    return send_file(output, as_attachment=True, download_name="talabalar_royxati.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.route("/admin/export/zip")
@admin_required
def export_zip():
    with db() as con:
        rows = con.execute(
            "SELECT * FROM students ORDER BY group_name, surname COLLATE NOCASE, first_name COLLATE NOCASE"
        ).fetchall()
    if not rows:
        flash("Eksport qilish uchun hali ma’lumot yo‘q.", "error")
        return redirect(url_for("admin"))

    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as zf:
        excel = make_excel(rows)
        zf.writestr("talabalar_royxati.xlsx", excel.getvalue())
        missing_images = []
        for row in rows:
            relative = Path(row["image_filename"])
            # Yo'l faqat saqlangan guruh/fayl ko'rinishida bo'lishi mumkin.
            if relative.is_absolute() or ".." in relative.parts:
                missing_images.append(row["image_filename"])
                continue
            path = UPLOAD_DIR / relative
            if path.is_file():
                group = secure_filename(row["group_name"])
                zf.write(path, f"Rasmlar/{group}/{path.name}")
            else:
                missing_images.append(row["image_filename"])
        if missing_images:
            zf.writestr("rasm_topilmadi.txt",
                        "Quyidagi rasm fayllari serverda topilmadi:\n" + "\n".join(missing_images))
    archive.seek(0)
    name = "talabalar_arxivi_" + datetime.now().strftime("%Y%m%d_%H%M") + ".zip"
    return send_file(archive, as_attachment=True, download_name=name, mimetype="application/zip")

@app.errorhandler(413)
def too_large(_error):
    return render_template("index.html", groups=GROUPS,
                           error_message="Yuklangan fayl juda katta. 7 MB dan kichik rasm tanlang."), 413

@app.route("/health")
def health():
    return {"status": "ok"}

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "5000")), debug=False)
